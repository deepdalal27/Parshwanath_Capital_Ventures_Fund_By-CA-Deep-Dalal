"""Workbook 02 - Anchor & QIB economics.

Issue structure -> anchor pool -> our likely allotment -> 30/90-day lock-in exits
by scenario, compared with a non-anchor QIB (ASBA) bid exited on listing day,
plus the supply-overhang calendar.

All sample values are ILLUSTRATIVE. Replace every blue cell from the offer document.
"""
import datetime as dt

from openpyxl import Workbook

from house import (pick, sample_banner, CR, DATE, DISCLAIMER, INT, PCT, F_BOLD, F_NOTE, InputBlock, dropdown, flag_text,
                   footer, header, note, put, title, widths)

D = dt.date


ANCHOR_POLICY = {"board", "qib_pct", "nii_pct", "rii_pct", "anc_pct", "mf_anc", "ins_anc", "mf_net", "min_anc",
                 "bid", "allot_pct", "qib_bid", "qib_sub", "blocked_days", "lock1_pct", "lock1_d", "lock2_d",
                 "cof", "txn", "exit"}


def build(path, facts=None):
    """facts=None builds the illustrative template; facts=dict (see FACTS_SCHEMA.md) builds an issuer."""
    wb = Workbook()

    # ------------------------------------------------------------------ Cover
    cv = wb.active
    cv.title = "Cover"
    title(cv, "Anchor & QIB Economics Calculator")
    lines = [
        "Purpose: size the anchor pool, estimate our discretionary allotment, model exits at the 30/90-day",
        "lock-in expiries, compare with a non-anchor QIB bid sold on listing day, and map supply overhang.",
        "",
        "How to use: 1) Fill every blue cell on 'Inputs' from the RHP (cite the page in column F).",
        "            2) Read 'Allocation' (pool, allotment, scenario returns, exit-route check).",
        "            3) Fill 'Overhang' rows from the RHP capital-structure / lock-in table.",
        "",
        "Colours: blue = input | black = formula | green = link from another sheet | yellow = key cell.",
        "Units: Rs crore unless stated. Share counts on 'Overhang' are in lakh.",
        "",
        "Anchor allocation is DISCRETIONARY (issuer + BRLM). A favourable call is not an allotment.",
        "Anchor shares are locked in: 50% for 30 days and 50% for 90 days from allotment.",
        "Returns are pre-tax. Tax character for a CAT III AIF depends on the PPM; refer out.",
        "",
        sample_banner(facts),
        "",
        DISCLAIMER,
    ]
    for i, t in enumerate(lines, start=4):
        cv[f"A{i}"] = t
    cv["A18"].font = F_BOLD
    cv["A20"].font = F_BOLD
    widths(cv, {"A": 110})

    # ----------------------------------------------------------------- Inputs
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs - from the offer document (blue cells only)")
    header(ws, 3, ["", "Item", "Value", "Unit", "Hint / page cite"], col=1)
    b = InputBlock(ws, 3, facts=None if facts is None else facts.get("inputs", {}), policy=ANCHOR_POLICY)
    b.section("Issue")
    b.add("name", "Issuer", "Illustrative Co. Ltd (REPLACE)", fmt=None, hint="Name as per RHP")
    b.add("board", "Platform", "Mainboard", fmt=None, hint="Mainboard / SME")
    dropdown(ws, b.addr["board"].split("!")[1].replace("$", ""), ["Mainboard", "SME"])
    b.add("offer", "Final offer price", 100, unit="Rs/share")
    b.add("anchor_px", "Anchor allocation price", 100, unit="Rs/share",
          hint="If offer price > anchor price, anchors pay the difference; if lower, no refund")
    b.add("issue", "Total issue size", 500, unit="Rs cr", hint="Fresh + OFS")
    b.add("emp", "Employee reservation", 5, unit="Rs cr")
    b.add("shr", "Shareholder reservation", 0, unit="Rs cr")
    b.section("Net-offer split (confirm from 'Offer Structure' chapter)")
    b.add("qib_pct", "QIB portion", 0.50, PCT, hint="Reg 6(1): not more than 50%; Reg 6(2): at least 75%")
    b.add("nii_pct", "Non-institutional portion", 0.15, PCT)
    b.add("rii_pct", "Retail portion", 0.35, PCT)
    b.add("anc_pct", "Anchor share of QIB portion", 0.60, PCT, hint="Maximum 60% of the QIB portion")
    b.add("mf_anc", "Anchor reserved: domestic MFs", 1 / 3, PCT, hint="One-third of anchor portion")
    b.add("ins_anc", "Anchor reserved: insurers & pension funds", 0.0667, PCT,
          hint="Per current ICDR amendment - CONFIRM split from the offer document")
    b.add("mf_net", "Net QIB reserved: MFs", 0.05, PCT, hint="5% of net QIB portion")
    b.section("Our bid")
    b.add("min_anc", "Minimum anchor application", 10, unit="Rs cr",
          hint="Mainboard Rs 10 cr - confirm platform limit in the offer document")
    b.add("bid", "Our anchor bid", 25, unit="Rs cr", key_cell=True)
    b.add("allot_pct", "Assumed allotment as % of bid", 0.60, PCT,
          hint="Discretionary - assumption, not entitlement")
    b.add("qib_bid", "Alternative: non-anchor QIB (ASBA) bid", 25, unit="Rs cr")
    b.add("qib_sub", "Expected QIB (ex-anchor) subscription", 50, '0.00"x"',
          hint="Allotment roughly proportionate: bid / subscription")
    b.add("blocked_days", "Days ASBA funds blocked", 6, INT, unit="days", hint="Issue open to listing")
    b.section("Dates and costs")
    b.add("bid_date", "Anchor bid / pay-in date", D(2026, 10, 12), DATE, hint="One working day before issue opens")
    b.add("allot_date", "Allotment date (basis of allotment)", D(2026, 10, 16), DATE,
          hint="Lock-in runs from this date")
    b.add("list_date", "Listing date", D(2026, 10, 19), DATE)
    b.add("lock1_pct", "Lock-in tranche 1 share", 0.50, PCT)
    b.add("lock1_d", "Lock-in tranche 1 days", 30, INT, unit="days")
    b.add("lock2_d", "Lock-in tranche 2 days", 90, INT, unit="days")
    b.add("cof", "Cost of funds / opportunity cost", 0.09, PCT, unit="p.a.",
          hint="Anchor money is PAID, not blocked")
    b.add("txn", "Exit transaction cost (brokerage+STT+other)", 0.0025, PCT, hint="On sale value")
    b.add("exit", "Planned exit route", "Lock-in expiry (30/90 days)", fmt=None, key_cell=True)
    dropdown(ws, b.addr["exit"].split("!")[1].replace("$", ""),
             ["Listing day", "Lock-in expiry (30/90 days)", "2-3 year hold"])

    # scenario price paths (move vs offer price)
    b.row += 1
    sr = b.row
    ws[f"B{sr}"] = "Price path vs offer price (illustrative)"
    ws[f"B{sr}"].font = F_BOLD
    header(ws, sr + 1, ["Scenario", "Listing day", "Day 30", "Day 90"], col=2)
    paths = {"Bear": (-0.10, -0.15, -0.20), "Neutral": (0.10, 0.08, 0.12), "Bull": (0.35, 0.30, 0.40)}
    if facts is not None:
        paths = {k: tuple(facts.get("paths", {}).get(k, (None, None, None))) for k in paths}
    sc = {}
    for i, (name, vals) in enumerate(paths.items()):
        r = sr + 2 + i
        put(ws, f"B{r}", name, "bold")
        for j, v in enumerate(vals):
            col = "CDE"[j]
            put(ws, f"{col}{r}", v, "input", PCT)
        sc[name] = r
    note(ws, f"B{sr + 5}", "Grey market premium is market colour only - never use it to move fair value.")
    widths(ws, {"A": 2, "B": 44, "C": 26, "D": 12, "E": 14, "F": 70})
    footer(ws)
    I = b.addr

    # ------------------------------------------------------------- Allocation
    al = wb.create_sheet("Allocation")
    title(al, "Anchor pool, our allotment, and exit returns")
    a = InputBlock(al, 3)
    a.section("Issue structure (Rs cr)")
    a.add("net", "Net offer", f"={I['issue']}-{I['emp']}-{I['shr']}", CR, "calc")
    a.add("qib", "QIB portion", f"={a.addr['net']}*{I['qib_pct']}", CR, "calc")
    a.add("nii", "Non-institutional portion", f"={a.addr['net']}*{I['nii_pct']}", CR, "calc")
    a.add("rii", "Retail portion", f"={a.addr['net']}*{I['rii_pct']}", CR, "calc")
    a.add("split_chk", "Split check (must total 100%)",
          f'=IF(ABS({I["qib_pct"]}+{I["nii_pct"]}+{I["rii_pct"]}-1)<0.0001,"OK","RED: portions do not total 100%")',
          None, "calc")
    a.add("anchor", "Anchor portion", f"={a.addr['qib']}*{I['anc_pct']}", CR, "calc", key_cell=True)
    a.add("anc_chk", "Anchor cap check", f'=IF({I["anc_pct"]}<=0.6,"OK","RED: anchor above 60% of QIB")', None, "calc")
    a.add("anc_mf", "  of which reserved for domestic MFs", f"={a.addr['anchor']}*{I['mf_anc']}", CR, "calc")
    a.add("anc_ins", "  of which reserved for insurers & pension funds", f"={a.addr['anchor']}*{I['ins_anc']}", CR, "calc")
    a.add("anc_open", "  open anchor pool (where an AIF competes)",
          f"={a.addr['anchor']}-{a.addr['anc_mf']}-{a.addr['anc_ins']}", CR, "calc", key_cell=True)
    a.add("qib_net", "Net QIB (after anchor)", f"={a.addr['qib']}-{a.addr['anchor']}", CR, "calc")
    a.add("qib_net_mf", "  of which reserved for MFs", f"={a.addr['qib_net']}*{I['mf_net']}", CR, "calc")
    a.add("anc_shares", "Anchor book size (lakh shares)", f"={a.addr['anchor']}*100/{I['anchor_px']}", CR, "calc")

    a.section("Our anchor allotment")
    a.add("bid_share", "Our bid as % of open anchor pool", f"={I['bid']}/{a.addr['anc_open']}", PCT, "calc")
    a.add("bid_chk", "Bid size check",
          f'=IF({I["bid"]}<{I["min_anc"]},"RED: below minimum anchor application",'
          f'IF({I["bid"]}>{a.addr["anc_open"]},"RED: bid exceeds open anchor pool","OK"))', None, "calc")
    a.add("allot_amt", "Assumed allotment (Rs cr)", f"=MIN({I['bid']}*{I['allot_pct']},{a.addr['anc_open']})", CR, "calc")
    a.add("eff_px", "Effective cost per share", f"=MAX({I['anchor_px']},{I['offer']})", CR, "calc", unit="Rs",
          hint="Anchors top up if offer price exceeds anchor price")
    a.add("shares", "Shares allotted", f"=ROUNDDOWN({a.addr['allot_amt']}*10^7/{I['anchor_px']},0)", INT, "calc")
    a.add("cost", "Total cost (Rs cr)", f"={a.addr['shares']}*{a.addr['eff_px']}/10^7", CR, "calc")
    a.add("t1", "Tranche 1 shares", f"=ROUNDDOWN({a.addr['shares']}*{I['lock1_pct']},0)", INT, "calc")
    a.add("t2", "Tranche 2 shares", f"={a.addr['shares']}-{a.addr['t1']}", INT, "calc")
    a.add("t1_date", "Tranche 1 free date", f"={I['allot_date']}+{I['lock1_d']}", DATE, "calc")
    a.add("t2_date", "Tranche 2 free date", f"={I['allot_date']}+{I['lock2_d']}", DATE, "calc")
    a.add("d1", "Days funds deployed - tranche 1", f"={a.addr['t1_date']}-{I['bid_date']}", INT, "calc")
    a.add("d2", "Days funds deployed - tranche 2", f"={a.addr['t2_date']}-{I['bid_date']}", INT, "calc")
    a.add("fund", "Funding cost (Rs cr)",
          f"=({a.addr['t1']}*{a.addr['d1']}+{a.addr['t2']}*{a.addr['d2']})*{a.addr['eff_px']}/10^7*{I['cof']}/365",
          CR, "calc")
    a.add("wdays", "Weighted days deployed",
          f"=({a.addr['t1']}*{a.addr['d1']}+{a.addr['t2']}*{a.addr['d2']})/{a.addr['shares']}", '0.0', "calc")

    a.section("Exit-route consistency")
    a.add("exit_chk", "Anchor bid vs planned exit",
          f'=IF({I["exit"]}="Listing day","INCONSISTENT: anchor shares are locked in ("&TEXT({I["lock1_pct"]},"0%")'
          f'&" for "&{I["lock1_d"]}&" days, rest for "&{I["lock2_d"]}&" days) and cannot be sold on listing day. '
          f'Bid as non-anchor QIB or change the exit plan.","OK: exit plan is compatible with anchor lock-in")',
          None, "calc", key_cell=True)
    a.add("disc", "Allotment note", "Anchor allotment is discretionary; there is no proportionate entitlement.",
          None, "calc")

    # scenario table
    r0 = a.row + 1
    al[f"B{r0}"] = "Scenario returns - ANCHOR route (exit at lock-in expiry)"
    al[f"B{r0}"].font = F_BOLD
    header(al, r0 + 1, ["Line", "Bear", "Neutral", "Bull"], col=2)
    rows = [
        ("lp", "Listing-day price (Rs)", lambda s: f"={I['offer']}*(1+Inputs!$C${sc[s]})", CR),
        ("p30", "Day-30 price (Rs)", lambda s: f"={I['offer']}*(1+Inputs!$D${sc[s]})", CR),
        ("p90", "Day-90 price (Rs)", lambda s: f"={I['offer']}*(1+Inputs!$E${sc[s]})", CR),
        ("mtm", "Mark-to-market on listing day, locked (Rs cr)",
         lambda s: f"={a.addr['shares']}*{{lp}}/10^7-{a.addr['cost']}", CR),
        ("g1", "Tranche 1 sale value (Rs cr)", lambda s: f"={a.addr['t1']}*{{p30}}/10^7", CR),
        ("g2", "Tranche 2 sale value (Rs cr)", lambda s: f"={a.addr['t2']}*{{p90}}/10^7", CR),
        ("gross", "Gross sale value (Rs cr)", lambda s: "={g1}+{g2}", CR),
        ("tc", "Transaction cost (Rs cr)", lambda s: f"={{gross}}*{I['txn']}", CR),
        ("pnl", "Net P&L after funding cost (Rs cr)",
         lambda s: f"={{gross}}-{{tc}}-{a.addr['cost']}-{a.addr['fund']}", CR),
        ("ret", "Net return on cost", lambda s: f"={{pnl}}/{a.addr['cost']}", PCT),
        ("ann", "Annualised (simple, 365-day)", lambda s: f"={{ret}}*365/{a.addr['wdays']}", PCT),
    ]
    out = {}
    for i, (k, lab, fn, fmt) in enumerate(rows):
        r = r0 + 2 + i
        al[f"B{r}"] = lab
        for j, s in enumerate(("Bear", "Neutral", "Bull")):
            col = "CDE"[j]
            f = fn(s).format(**{kk: f"{col}{rr}" for kk, rr in out.items()})
            put(al, f"{col}{r}", f, "calc", fmt)
        out[k] = r
    anchor_rows = dict(out)

    # QIB ASBA comparison
    r1 = r0 + 2 + len(rows) + 1
    al[f"B{r1}"] = "Comparison - NON-ANCHOR QIB route (ASBA, sold on listing day)"
    al[f"B{r1}"].font = F_BOLD
    header(al, r1 + 1, ["Line", "Bear", "Neutral", "Bull"], col=2)
    q_amt = f"MIN({I['qib_bid']},{I['qib_bid']}/{I['qib_sub']})"
    q_sh = f"ROUNDDOWN({q_amt}*10^7/{I['offer']},0)"
    qrows = [
        ("qsh", "Shares allotted (proportionate estimate)", lambda c: f"={q_sh}", INT),
        ("qcost", "Cost (Rs cr)", lambda c: f"={{qsh}}*{I['offer']}/10^7", CR),
        ("qsale", "Listing-day sale value (Rs cr)", lambda c: f"={{qsh}}*{c}{anchor_rows['lp']}/10^7", CR),
        ("qblock", "Opportunity cost of blocked funds (Rs cr)",
         lambda c: f"={I['qib_bid']}*{I['cof']}*{I['blocked_days']}/365", CR),
        ("qpnl", "Net P&L (Rs cr)", lambda c: f"={{qsale}}*(1-{I['txn']})-{{qcost}}-{{qblock}}", CR),
        ("qret", "Net return on allotted cost", lambda c: f"=IFERROR({{qpnl}}/{{qcost}},0)", PCT),
        ("diff", "Anchor P&L minus QIB P&L (Rs cr)", lambda c: f"={c}{anchor_rows['pnl']}-{{qpnl}}", CR),
    ]
    qout = {}
    for i, (k, lab, fn, fmt) in enumerate(qrows):
        r = r1 + 2 + i
        al[f"B{r}"] = lab
        for j in range(3):
            col = "CDE"[j]
            f = fn(col).format(**{kk: f"{col}{rr}" for kk, rr in qout.items()})
            put(al, f"{col}{r}", f, "calc", fmt)
        qout[k] = r
    rn = r1 + 2 + len(qrows) + 1
    note(al, f"B{rn}", "Anchor route earns on a larger, more certain allotment but carries 30/90-day price risk "
                       "and pays funds upfront; QIB route is small and allotment-uncertain but liquid on day one.")
    flag_text(al, "C1:C200")
    widths(al, {"A": 2, "B": 50, "C": 30, "D": 16, "E": 16, "F": 50})
    footer(al)

    # --------------------------------------------------------------- Overhang
    ov = wb.create_sheet("Overhang")
    title(ov, "Supply-overhang calendar (shares in lakh)")
    put(ov, "B4", "Total post-issue shares (lakh)", "bold")
    put(ov, "C4", pick(facts, "total_shares_lakh", 1000, None), "input", CR)
    put(ov, "B5", "Free float at listing (lakh)", "bold")
    put(ov, "B6", "Free float as % of total", "bold")
    header(ov, 8, ["Category", "Shares (lakh)", "Lock-in days from allotment", "Release date",
                   "Free float just before release", "Release as % of float", "Cumulative released",
                   "Flag", "Page cite / note"], col=2)
    lock_rows = [
        ("Anchor - tranche 1", f"={a.addr['anc_shares']}*{I['lock1_pct']}", f"={I['lock1_d']}", "link",
         "Whole anchor book, not only our allotment"),
        ("Anchor - tranche 2", f"={a.addr['anc_shares']}*(1-{I['lock1_pct']})", f"={I['lock2_d']}", "link", ""),
        ("Pre-issue non-promoter shareholders", 150, 180, "input", "6 months - confirm"),
        ("Promoter holding above minimum contribution", 300, 180, "input", "Confirm period in RHP lock-in table"),
        ("Promoter minimum contribution", 200, 540, "input", "18 months / 3 years if capex objects - confirm"),
        ("Market maker inventory (SME only)", 0, 0, "input", "SME: market maker is the liquidity"),
        ("Other (specify)", 0, 0, "input", ""),
    ]
    if facts is not None:  # issuer: anchor link rows + rows from the RHP lock-in table
        given = [(r[0], r[1], r[2], "input", r[3] if len(r) > 3 else "") for r in facts.get("overhang", [])]
        lock_rows = lock_rows[:2] + given + [("", None, None, "input", "")] * max(0, 5 - len(given))
    first, last = 9, 9 + len(lock_rows) - 1
    for i, (cat, sh, days, kind, hint) in enumerate(lock_rows):
        r = first + i
        put(ov, f"B{r}", cat, "bold")
        put(ov, f"C{r}", sh, kind, CR)
        put(ov, f"D{r}", days, kind, INT)
        put(ov, f"E{r}", f"={I['allot_date']}+D{r}", "calc", DATE)
        put(ov, f"F{r}", f'=$C$5+SUMIF($E${first}:$E${last},"<"&E{r},$C${first}:$C${last})', "calc", CR)
        put(ov, f"G{r}", f"=IF(F{r}>0,C{r}/F{r},0)", "calc", PCT)
        put(ov, f"H{r}", f'=SUMIF($E${first}:$E${last},"<="&E{r},$C${first}:$C${last})', "calc", CR)
        put(ov, f"I{r}", f'=IF(C{r}=0,"-",IF(G{r}>0.25,"RED: large cliff vs float",IF(G{r}>0.1,"AMBER: watch","OK")))',
            "calc")
        note(ov, f"J{r}", hint)
    put(ov, "C5", f"=C4-SUM(C{first}:C{last})", "calc", CR, key=True)
    put(ov, "C6", "=C5/C4", "calc", PCT)
    note(ov, f"B{last + 2}", "A small float meeting a large release is a price event unrelated to the business. "
                             "Re-key rows from the RHP 'Capital Structure' lock-in table.")
    flag_text(ov, f"I{first}:I{last}")
    widths(ov, {"A": 2, "B": 42, "C": 14, "D": 14, "E": 14, "F": 16, "G": 14, "H": 14, "I": 26, "J": 46})
    footer(ov)

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return {"inputs": I, "alloc": a.addr, "anchor_rows": anchor_rows, "qib_rows": qout, "scen_rows": sc,
            "overhang_rows": (first, last)}


if __name__ == "__main__":
    build("02_Anchor_QIB_Economics.xlsx")
