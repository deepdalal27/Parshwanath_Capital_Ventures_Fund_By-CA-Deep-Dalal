"""Workbook 01 - DRHP / RHP screening checklist.

Basis & staleness banner, quick forensic metrics with auto-flags, the ten forensic
checks plus regulatory checks, the twelve analyst work blocks, a data-gaps register
and a scorecard. All sample values are ILLUSTRATIVE.
"""
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment

from house import (finalize, sample_banner, CR, DATE, DISCLAIMER, F_BOLD, INT, PCT, X, WRAP, InputBlock, dropdown, flag_text, footer,
                   header, note, put, title, widths)

cell = lambda a: a.split("!")[1].replace("$", "")
STATUS = ["PASS", "WATCH", "FAIL", "N.A.", "OPEN"]


SCREEN_POLICY = {"platform", "stage", "band", "bid_cat", "exit"}


def build(path, facts=None):
    """facts=None builds the illustrative template; facts=dict (see FACTS_SCHEMA.md) builds an issuer."""
    wb = Workbook()
    cv = wb.active
    cv.title = "Cover"
    title(cv, "DRHP / RHP Screening Checklist")
    for i, t in enumerate([
        "A DRHP is a screen; an RHP is a decision. Record the stage on 'Basis' first.",
        "Order: Basis > Quick Metrics (auto-flags) > Forensic Checks > Work Blocks > Data Gaps > Scorecard.",
        "Two rules: never invent a figure (leave blank, log it in Data Gaps); never draft disclaimer text.",
        "Every figure as restated, cited to the printed page. Watch for three page-numbering blocks.",
        "If the restated-financials annexure is image-only, OCR to locate but read the image to cite.",
        "Thresholds in grey cells on 'Quick Metrics' are house defaults and editable.",
        "", "Colours: blue = input | black = formula | yellow = key cell.",
        sample_banner(facts), "", DISCLAIMER]):
        cv[f"A{i + 4}"] = t
    cv["A12"].font = F_BOLD
    cv["A14"].font = F_BOLD
    widths(cv, {"A": 110})

    # ------------------------------------------------------------------ Basis
    bs = wb.create_sheet("Basis")
    title(bs, "Basis and staleness banner")
    b = InputBlock(bs, 3, facts=None if facts is None else facts.get("basis", {}), policy=SCREEN_POLICY)
    b.section("Document")
    b.add("name", "Issuer", "Illustrative Co. Ltd (REPLACE)", None)
    b.add("platform", "Platform", "Mainboard", None)
    dropdown(bs, cell(b.addr["platform"]), ["Mainboard", "SME"])
    b.add("stage", "Stage", "DRHP", None, key_cell=True)
    dropdown(bs, cell(b.addr["stage"]), ["DRHP", "RHP", "Addendum", "Prospectus"])
    b.add("doc_date", "Document date", dt.date(2026, 8, 20), DATE)
    b.add("aud_to", "Audited financials up to", dt.date(2026, 3, 31), DATE)
    b.add("today", "Today", "=TODAY()", DATE, "calc")
    b.add("stale", "Staleness (days since audited period end)", f"={b.addr['today']}-{b.addr['aud_to']}", INT, "calc")
    b.add("stale_flag", "Staleness flag",
          f'=IF({b.addr["stale"]}>365,"RED: audited record over a year old - commit no capital until the RHP is read",'
          f'IF({b.addr["stale"]}>180,"AMBER: RHP must carry financials no more than six months stale","OK"))',
          None, "calc")
    b.add("band", "Price band announced?", "No", None)
    dropdown(bs, cell(b.addr["band"]), ["Yes", "No"])
    b.add("brlm", "Book running lead manager(s)", "[name]", None, hint="Compile BRLM track record (CIR/MIRSD/1/2012)")
    b.add("outside", "Figures from outside the offer document", "[source + date, or 'none']", None)
    b.section("Trade being underwritten")
    b.add("bid_cat", "Bid category", "Anchor", None)
    dropdown(bs, cell(b.addr["bid_cat"]), ["Anchor", "QIB (non-anchor)", "NII"])
    b.add("exit", "Exit route underwritten", "2-3 year hold (assumed)", None, key_cell=True,
          hint="Ask per issuer; default is the 2-3 year hold")
    dropdown(bs, cell(b.addr["exit"]), ["2-3 year hold (assumed)", "2-3 year hold (confirmed)", "Lock-in expiry (30/90 days)", "Listing day"])
    b.add("exit_chk", "Bid category vs exit route",
          f'=IF(AND({b.addr["bid_cat"]}="Anchor",{b.addr["exit"]}="Listing day"),'
          f'"INCONSISTENT: anchor allotment is locked in 30/90 days - cannot exit on listing day","OK")', None, "calc")
    b.add("tool", "Operative valuation tool",
          f'=IF({b.addr["exit"]}="Listing day","Listing-day framework (DCF used only as a floor test)",'
          f'"DCF governs (intrinsic value over the hold)")', None, "calc")
    flag_text(bs, "C1:C40")
    widths(bs, {"A": 2, "B": 44, "C": 44, "D": 8, "E": 50})
    footer(bs)
    B = b.addr

    # ----------------------------------------------------------- Quick Metrics
    qm = wb.create_sheet("Quick Metrics")
    title(qm, "Quick forensic metrics (Rs cr) - restated, as reported; stub never annualised")
    header(qm, 4, ["Line item", "FY-3", "FY-2", "FY-1 (latest)", "Stub (as reported)", "Page cite"], col=2)
    items = [
        ("rev", "Revenue from operations", (620, 850, 1000, 540)),
        ("oth", "Other income", (8, 10, 22, 6)),
        ("pbt", "Profit before tax", (60, 85, 108, 58)),
        ("ctax", "Current tax", (14, 20, 24, 13)),
        ("dtax", "Deferred tax", (1, 2, 4, 2)),
        ("pat", "PAT", (45, 63, 80, 43)),
        ("cfo", "Cash flow from operations", (30, 28, 52, 15)),
        ("rpt_s", "Related-party sales", (12, 30, 55, 20)),
        ("rpt_p", "Related-party purchases (incl. capex)", (5, 6, 9, 4)),
        ("nw", "Net worth", (300, 360, 450, 490)),
        ("recv", "Trade receivables", (110, 160, 210, 230)),
        ("inv", "Inventory", (90, 120, 140, 150)),
        ("pay", "Trade payables", (70, 115, 150, 160)),
        ("cl", "Contingent liabilities", (20, 35, 60, 60)),
        ("gw", "Goodwill + intangibles", (0, 0, 12, 12)),
        ("reval", "Revaluation reserve", (0, 0, 0, 0)),
    ]
    R = {}
    for i, (k, lab, vals) in enumerate(items):
        r = 5 + i
        if facts is not None:
            vals = tuple(facts.get("items", {}).get(k, (None,) * 4))
        cite = "" if facts is None else facts.get("cites", {}).get(k, "")
        qm[f"B{r}"] = lab
        for j, v in enumerate(vals):
            put(qm, f"{'CDEF'[j]}{r}", v, "input", CR)
        put(qm, f"G{r}", cite, "input")
        R[k] = r
    r = 5 + len(items) + 1
    qm[f"B{r}"] = "Single-period inputs (latest)"
    qm[f"B{r}"].font = F_BOLD
    single = [("top1c", "Top customer % of revenue", 0.22, PCT), ("top10c", "Top-10 customers % of revenue", 0.68, PCT),
              ("top1s", "Top supplier % of purchases", 0.35, PCT), ("geo", "Largest state/country % of revenue", 0.55, PCT),
              ("fresh_net", "Fresh issue net of expenses", 190, CR), ("waca", "Promoter weighted avg cost of acquisition (Rs)", 12, CR),
              ("offer", "Offer price / cap (Rs)", 190, CR)]
    for k, lab, v, fmt in single:
        r += 1
        if facts is not None:
            v = facts.get("single", {}).get(k)
        qm[f"B{r}"] = lab
        put(qm, f"E{r}", v, "input", fmt)
        R[k] = r
    L = lambda k: f"$E${R[k]}"  # latest-year column
    rng = lambda k: f"$C${R[k]}:$E${R[k]}"  # three full years (stub excluded from cumulative)

    r += 2
    m0 = r
    qm[f"B{r}"] = "Auto metrics"
    qm[f"B{r}"].font = F_BOLD
    header(qm, r + 1, ["Metric", "Value", "Threshold", "Flag", "Forensic check / block"], col=2)
    metrics = [
        ("Cumulative CFO / cumulative PAT (3 yrs)", f"=SUM({rng('cfo')})/SUM({rng('pat')})", PCT, 0.70, "lt",
         "Check 2 / Block 1"),
        ("Cumulative cash tax rate (current tax / PBT)", f"=SUM({rng('ctax')})/SUM({rng('pbt')})", PCT, 0.15, "lt",
         "Check 3"),
        ("Deferred tax share of total tax (latest)", f"={L('dtax')}/({L('ctax')}+{L('dtax')})", PCT, 0.50, "gt", "Check 3"),
        ("Other income % of PBT (latest)", f"={L('oth')}/{L('pbt')}", PCT, 0.15, "gt", "Block 1"),
        ("RPT sales % of revenue (latest)", f"={L('rpt_s')}/{L('rev')}", PCT, 0.02, "gt", "Block 9 (2% flag)"),
        ("RPT purchases % of net worth (latest)", f"={L('rpt_p')}/{L('nw')}", PCT, 0.02, "gt", "Check 4 / Block 9"),
        ("RPT share of revenue growth (FY-2 to latest)",
         f"=IFERROR(({L('rpt_s')}-$D${R['rpt_s']})/({L('rev')}-$D${R['rev']}),\"\")", PCT, 0.10, "gt", "Block 4"),
        ("Top customer concentration", f"={L('top1c')}", PCT, 0.50, "gt", "Check 8"),
        ("Top-10 customer concentration", f"={L('top10c')}", PCT, 0.80, "gt", "Check 8"),
        ("Top supplier concentration", f"={L('top1s')}", PCT, 0.50, "gt", "Check 8 (supply side)"),
        ("Geographic concentration", f"={L('geo')}", PCT, 0.80, "gt", "Check 8"),
        ("Contingent liabilities / net worth", f"={L('cl')}/{L('nw')}", PCT, 0.25, "gt", "Check 9"),
        ("Goodwill + intangibles / net worth", f"={L('gw')}/{L('nw')}", PCT, 0.30, "gt", "Check 7"),
        ("Revaluation reserve (any non-zero fails)", f"={L('reval')}", CR, 0, "gt", "Check 6"),
        ("Receivable days (closing, latest)", f"={L('recv')}/{L('rev')}*365", '0.0', 90, "gt", "Block 3"),
        ("Inventory days (closing, on revenue)", f"={L('inv')}/{L('rev')}*365", '0.0', 90, "gt", "Block 3"),
        ("Trade NWC % of revenue (latest)", f"=({L('recv')}+{L('inv')}-{L('pay')})/{L('rev')}", PCT, 0.25, "gt", "Block 2"),
        ("Incremental trade NWC per Rs of incremental revenue",
         f"=(({L('recv')}+{L('inv')}-{L('pay')})-($D${R['recv']}+$D${R['inv']}-$D${R['pay']}))/({L('rev')}-$D${R['rev']})",
         '0.00', 0.30, "gt", "Block 2"),
        ("ROE pre-issue (latest)", f"={L('pat')}/{L('nw')}", PCT, None, None, "Block 8"),
        ("ROE post-issue (same PAT)", f"={L('pat')}/({L('nw')}+{L('fresh_net')})", PCT, None, None, "Block 8"),
        ("Offer price / promoter WACA", f"={L('offer')}/{L('waca')}", X, 10, "gt", "Block 8 / 11"),
    ]
    for i, (lab, f, fmt, thr, op, ref) in enumerate(metrics):
        rr = m0 + 2 + i
        qm[f"B{rr}"] = lab
        # A metric built on a blank single-period input is a data gap, not a zero: show it blank.
        need = sorted(set(re.findall(r"\$E\$(\d+)", f)))
        if need:
            f = '=IF(OR(' + ",".join(f"ISBLANK($E${n})" for n in need) + '),"",' + f[1:] + ")"
        put(qm, f"C{rr}", f, "calc", fmt)
        if thr is not None:
            put(qm, f"D{rr}", thr, "input", fmt)
            cmp_ = f"C{rr}<D{rr}" if op == "lt" else f"C{rr}>D{rr}"
            put(qm, f"E{rr}", f'=IF(NOT(ISNUMBER(C{rr})),"-",IF({cmp_},"WATCH","OK"))', "calc")
        else:
            put(qm, f"E{rr}", "info", "calc")
        qm[f"F{rr}"] = ref
    flag_text(qm, f"E{m0 + 2}:E{m0 + 2 + len(metrics)}")
    note(qm, f"B{m0 + 3 + len(metrics)}", "CFO falling while PAT rises is the signal. Cumulative cash tax at the statutory "
                                          "rate is the strongest positive fact available - say it as plainly as the negatives.")
    widths(qm, {"A": 2, "B": 50, "C": 14, "D": 14, "E": 16, "F": 22, "G": 14})
    footer(qm)

    # -------------------------------------------------------- Forensic Checks
    fc = wb.create_sheet("Forensic Checks")
    title(fc, "Ten forensic checks + regulatory checks")
    header(fc, 4, ["#", "Check", "What to test", "Where in the document", "Status", "Page cite", "Finding"], col=1)
    checks = [
        ("1", "Bonus issue before filing", "Basis for Issue Price on pre-bonus count? EPS/NAV restated for all periods "
         "(Ind AS 33)? Consistent share count in every chapter?", "Capital Structure; Basis for Offer Price; Financials"),
        ("2", "Cash flow statement and CFO vs PAT", "Statement present? Cumulative CFO / PAT, standalone and consolidated "
         "(see Quick Metrics).", "Restated cash flow statement"),
        ("3", "Current vs deferred tax", "Near-nil current tax? Cumulative cash tax rate vs statutory.", "Restated P&L, tax note"),
        ("4", "Related-party capex and proceeds into subsidiaries", "Capitalised RP purchases; instrument, rate, tenor, "
         "security of funds to subsidiaries; leakage to minorities.", "RPT note; Objects"),
        ("5", "Objects arithmetic as price floor", "Objects / fixed fresh shares, grossed up for expenses and GCP. "
         "Floor above fair value?", "Objects of the Offer"),
        ("6", "Revaluation reserves", "Any non-zero balance.", "Other equity note"),
        ("7", "Goodwill vs tangible equity", "Goodwill share of equity; purchase accounting changed between documents?",
         "Restated balance sheet"),
        ("8", "Concentration - demand AND supply side", "Customer, supplier, geography; >80% flags; long-term supply "
         "agreement exists?", "Business; Risk Factors"),
        ("9", "Contingent liabilities and materiality policy", "CL vs net worth; is the materiality threshold set to "
         "exclude real matters?", "Outstanding Litigation; CL note"),
        ("10", "Investor presentation vs filed document", "Base years, margin definitions, customers, capacity, awards "
         "absent from the DRHP; KPIs outside the KPI table.", "Investor deck vs DRHP"),
        ("R1", "Auditor qualifications / emphasis of matter", "Any qualification, EoM, or CARO adverse remark.", "Auditor reports"),
        ("R2", "CARO 2020 / Rule 11(g) audit trail", "Audit-trail compliance reported?", "Auditor report annexure"),
        ("R3", "Monitoring agency (ICDR Reg 262(1))", "Appointed where required by issue size?", "Objects"),
        ("R4", "Mainboard eligibility (Reg 6(1)/6(2))", "Re-perform the eligibility computation; compare to KPI basis.",
         "Other Regulatory and Statutory Disclosures"),
        ("R5", "SME (Chapter IX) - post-2025 conditions", "OFS cap, selling-holder cap, GCP cap, operating-profit test, "
         "no proceeds to repay promoter/RP loans. Confirm against current ICDR text.", "Offer structure; Objects"),
        ("R6", "BRLM track record (CIR/MIRSD/1/2012)", "Compile from exchange data: issue price, listing gain, current "
         "price vs issue price.", "BRLM website / exchange data"),
        ("R7", "Firm means of finance (Reg 230(1)(e))", "75% of stated means of finance tied up, excluding issue proceeds?",
         "Objects"),
    ]
    for i, (n, ck, what, where) in enumerate(checks):
        r = 5 + i
        for j, v in enumerate((n, ck, what, where)):
            c = fc.cell(row=r, column=1 + j, value=v)
            c.alignment = WRAP
        got = {} if facts is None else facts.get("checks", {}).get(n, {})
        put(fc, f"E{r}", got.get("status", "OPEN"), "input")
        dropdown(fc, f"E{r}", STATUS)
        put(fc, f"F{r}", got.get("page", ""), "input")
        put(fc, f"G{r}", got.get("finding", ""), "input")
        fc[f"G{r}"].alignment = WRAP
        fc.row_dimensions[r].height = 42
    flag_text(fc, f"E5:E{4 + len(checks)}")
    widths(fc, {"A": 5, "B": 34, "C": 60, "D": 34, "E": 10, "F": 12, "G": 50})
    footer(fc)
    n_checks = len(checks)

    # ------------------------------------------------------------ Work Blocks
    wk = wb.create_sheet("Work Blocks")
    title(wk, "Twelve analyst work blocks - work all twelve, report what moves the conclusion")
    header(wk, 4, ["Block", "Question", "Finding", "Moves conclusion?", "Page cite"], col=1)
    blocks = {
        "1 Earnings quality": ["Cumulative CFO / PAT; CFO before and after working capital",
                               "Non-recurring items in other income as % of PBT",
                               "Is the stub period seasonally favourable for this industry?"],
        "2 Working capital": ["NWC % revenue each period", "Incremental NWC per Rs of incremental revenue",
                              "Closing-balance days vs issuer's average-based days"],
        "3 Receivables / payables / inventory": ["Schedule III ageing: buckets from due date or invoice date?",
                                                 "Provision coverage on old debt", "MSME payables and Sec 15/16 interest",
                                                 "Slow-moving inventory"],
        "4 Revenue integrity": ["Standalone vs consolidated revenue and EPS", "Related-party contribution to growth",
                                "Segment/product splits ledgered or derived (round numbers + one plug)?"],
        "5 Margin bridge": ["Gross margin / employee / other cost contribution to margin change",
                            "Is the margin plausible for the business model?", "Costs the company can reinstate at will"],
        "6 Balance sheet and banking": ["Sanctioned vs drawn, on-demand share", "Covenants and lender consents",
                                        "Personal guarantees / promoter property released by proceeds?",
                                        "Credit rating trend; insurance cover vs assets"],
        "7 Objects tested": ["% to productive assets / WC / debt / unallocated", "Debt repayment that is revolving WC",
                             "Quotation validity, FX date, firm orders, who prepared the DPR"],
        "8 Post-issue dilution": ["Post-money ROE / ROCE", "Offer price vs post-issue NAV (book dilution)",
                                  "Promoter WACA vs offer price"],
        "9 Related parties": ["Sales/purchases/advances/rent/remuneration/loans vs base",
                              "Promoter remuneration: contracted vs drawn (charge the gap in Bear)",
                              "Group entities that compete, supply or buy"],
        "10 Management and governance": ["Listed-company experience; KMP tenure < 1 year",
                                         "Independent director churn within 12 months",
                                         "Auditor forensics: changes, ADT-1/ADT-3 dates, resignation reason, firm size"],
        "11 Corporate record": ["Compounding / adjudication; late e-form filings", "BEN-1/BEN-2 compliance",
                                "Pre-IPO gifts/transfers and effect on WACA", "Statutory dues delays; approvals pending"],
        "12 Industry and competition": ["Who paid for the industry report?", "Competitors the document does not name",
                                        "Disposals: who bought and what they did next",
                                        "Arm's-length transactions in own shares - implied multiple"],
    }
    r = 5
    for blk, qs in blocks.items():
        for q in qs:
            wk[f"A{r}"] = blk
            wk[f"B{r}"] = q
            wk[f"B{r}"].alignment = WRAP
            put(wk, f"C{r}", "", "input")
            put(wk, f"D{r}", "", "input")
            dropdown(wk, f"D{r}", ["Yes", "No"])
            put(wk, f"E{r}", "", "input")
            r += 1
    widths(wk, {"A": 32, "B": 62, "C": 60, "D": 12, "E": 12})
    footer(wk)

    # --------------------------------------------------------------- Data Gaps
    dg = wb.create_sheet("Data Gaps")
    title(dg, "Data gaps - declared holes, never estimated")
    header(dg, 4, ["Item", "Severity", "Why it matters", "Source to try (BSE/NSE/MCA/ROC)", "Status", "Closed in RHP?"], col=1)
    for i in range(25):
        rr = 5 + i
        for c in "ABCDEF":
            put(dg, f"{c}{rr}", "", "input")
        dropdown(dg, f"B{rr}", ["High", "Medium", "Low"])
        dropdown(dg, f"E{rr}", ["Open", "Sourced", "Asked issuer/BRLM"])
        dropdown(dg, f"F{rr}", ["Closed", "Closed unsatisfactorily", "Ignored", "-"])
    put(dg, "A5", "BRLM track record compilation", "input")
    put(dg, "B5", "Medium", "input")
    put(dg, "C5", "Required by SEBI circular; usually delegated to the BRLM website", "input")
    widths(dg, {"A": 40, "B": 10, "C": 48, "D": 34, "E": 18, "F": 22})
    footer(dg)

    # --------------------------------------------------------------- Scorecard
    sc = wb.create_sheet("Scorecard")
    title(sc, "Forensic scorecard and call")
    put(sc, "B4", f'={B["name"]}&"  |  "&{B["stage"]}&"  |  "&{B["exit"]}', "calc")
    header(sc, 6, ["Status", "Count"], col=2)
    for i, s in enumerate(STATUS):
        r = 7 + i
        put(sc, f"B{r}", s, "bold")
        put(sc, f"C{r}", f"=COUNTIF('Forensic Checks'!$E$5:$E${4 + n_checks},B{r})", "calc", INT)
    put(sc, "B13", "Auto-metric WATCH flags", "bold")
    put(sc, "C13", f"=COUNTIF('Quick Metrics'!$E$1:$E$200,\"WATCH\")", "calc", INT)
    put(sc, "B14", "High-severity data gaps open", "bold")
    put(sc, "C14", "=COUNTIFS('Data Gaps'!$B$5:$B$29,\"High\",'Data Gaps'!$E$5:$E$29,\"Open\")", "calc", INT)
    put(sc, "B15", "Staleness", "bold")
    put(sc, "C15", f"={B['stale_flag']}", "link")
    put(sc, "B16", "Exit consistency", "bold")
    put(sc, "C16", f"={B['exit_chk']}", "link")
    flag_text(sc, "C15:C16")
    put(sc, "B18", "Screen outcome", "bold")
    put(sc, "C18", '=IF(C9>0,"FAIL items present - move to full DD only with explicit reasons",'
                   'IF(C11>0,"WATCH: "&C11&" checks still open - screen not complete",'
                   'IF(C8+C13>2,"WATCH: several watch items - full DD required","PASS to full due diligence")))',
        "calc", key=True)
    flag_text(sc, "C18")
    put(sc, "B20", "Call (Subscribe / Avoid / Neutral)", "bold")
    put(sc, "C20", "", "input", key=True)
    dropdown(sc, "C20", ["Subscribe", "Avoid", "Neutral"])
    sc["B22"] = "What would change our view (measurable, verifiable against the RHP)"
    sc["B22"].font = F_BOLD
    for i in range(5):
        put(sc, f"B{23 + i}", f"{i + 1}.", "bold")
        put(sc, f"C{23 + i}", "", "input")
    sc["B29"] = DISCLAIMER
    sc["B29"].font = F_BOLD
    widths(sc, {"A": 2, "B": 36, "C": 90})
    footer(sc)

    finalize(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return R, m0


if __name__ == "__main__":
    build("01_DRHP_Screening_Checklist.xlsx")
